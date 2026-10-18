"""Lectura del informe "Ingresos Operaciones por Concepto" exportado desde el SAC.

Acepta el libro Excel de gestión completo (busca la hoja con los encabezados del SAC)
o un export en CSV. Solo se leen las columnas crudas A–V; las columnas calculadas con
BUSCARV del libro se ignoran porque la clasificación la hace `core.clasificacion`.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Callable, Optional

import pandas as pd

from .esquema import COLUMNAS_OBLIGATORIAS, COLUMNAS_SAC

_EPOCA_EXCEL = datetime(1899, 12, 30)
_COLUMNAS_FECHA = ["FECHA_EMISION", "FECHA_VENCIMIENTO", "FECHA_PAGO", "PERIODO"]
_COLUMNAS_NUMERICAS = ["UNIDAD_PRESTACION_SERVICIO", "UNIDAD_FACTURACION", "MONTO", "MES", "TASA", "TIPO_DE_CAMBIO"]


@dataclass
class ResultadoLectura:
    datos: pd.DataFrame
    hoja: str
    filas_leidas: int
    advertencias: list[str] = field(default_factory=list)


def normalizar_encabezado(valor) -> str:
    """'UNIDAD_PRESTACION_SERVICIO_ ' -> 'UNIDAD_PRESTACION_SERVICIO'."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor)).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto.strip().upper())
    return texto.strip("_")


def _mapa_columnas(encabezados) -> dict[str, int]:
    mapa = {}
    for i, h in enumerate(encabezados):
        nombre = normalizar_encabezado(h)
        if nombre in COLUMNAS_SAC and nombre not in mapa:
            mapa[nombre] = i
    return mapa


def _es_encabezado(fila) -> bool:
    return all(c in _mapa_columnas(fila) for c in COLUMNAS_OBLIGATORIAS)


def _a_bytes(origen) -> bytes:
    if isinstance(origen, (bytes, bytearray)):
        return bytes(origen)
    if hasattr(origen, "getvalue"):
        return origen.getvalue()
    if hasattr(origen, "read"):
        return origen.read()
    with open(origen, "rb") as f:
        return f.read()


def leer_informe(origen, nombre: str = "", progreso: Optional[Callable[[int], None]] = None) -> ResultadoLectura:
    """Lee un informe SAC desde ruta, bytes o archivo subido (xlsx, xlsm o csv)."""
    nombre = (nombre or getattr(origen, "name", "") or str(origen if isinstance(origen, str) else "")).lower()
    contenido = _a_bytes(origen)
    if nombre.endswith((".csv", ".txt")) or not contenido[:2] == b"PK":
        return _leer_csv(contenido)
    return _leer_excel(contenido, progreso)


def _leer_excel(contenido: bytes, progreso) -> ResultadoLectura:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            encabezado, fila_enc = None, 0
            for n, fila in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
                if fila and _es_encabezado(fila):
                    encabezado, fila_enc = fila, n
                    break
            if encabezado is None:
                continue
            mapa = _mapa_columnas(encabezado)
            max_col = max(mapa.values()) + 1
            columnas = list(mapa.keys())
            indices = [mapa[c] for c in columnas]
            registros = []
            for n, fila in enumerate(ws.iter_rows(min_row=fila_enc + 1, max_col=max_col, values_only=True)):
                if fila is None or all(v is None for v in fila):
                    continue
                registros.append([fila[i] if i < len(fila) else None for i in indices])
                if progreso and n % 5000 == 0:
                    progreso(n)
            df = pd.DataFrame(registros, columns=columnas)
            return _normalizar(df, ws.title)
    finally:
        wb.close()
    raise ValueError(
        "No se encontró una hoja con los encabezados del informe SAC "
        f"(se requieren al menos: {', '.join(COLUMNAS_OBLIGATORIAS)})."
    )


def _leer_csv(contenido: bytes) -> ResultadoLectura:
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        separador = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t|").delimiter
    except csv.Error:
        separador = ";"
    df = pd.read_csv(io.StringIO(texto), sep=separador, dtype=str)
    df.columns = [normalizar_encabezado(c) for c in df.columns]
    faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in df.columns]
    if faltantes:
        raise ValueError(f"El CSV no tiene las columnas obligatorias: {', '.join(faltantes)}")
    df = df[[c for c in COLUMNAS_SAC if c in df.columns]]
    return _normalizar(df, "CSV")


def _a_numero(serie: pd.Series) -> pd.Series:
    if serie.dtype == object:
        texto = serie.astype(str).str.strip()
        # Formato chileno "1.234.567,89" -> "1234567.89"
        chileno = texto.str.contains(r",\d{1,2}$", regex=True) | texto.str.contains(r"\.\d{3}(?:\.|$)", regex=True)
        texto = texto.where(~chileno, texto.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
        return pd.to_numeric(texto.replace({"None": None, "nan": None, "": None}), errors="coerce")
    return pd.to_numeric(serie, errors="coerce")


def _a_fecha(serie: pd.Series) -> pd.Series:
    def convertir(v):
        if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
            return pd.NaT
        if isinstance(v, datetime):
            return v
        if isinstance(v, (int, float)):
            return _EPOCA_EXCEL + timedelta(days=float(v))
        texto = str(v).strip()
        if re.fullmatch(r"\d+(\.\d+)?", texto):
            return _EPOCA_EXCEL + timedelta(days=float(texto))
        return pd.to_datetime(texto, dayfirst=True, errors="coerce")

    return pd.to_datetime(serie.map(convertir), errors="coerce")


def _normalizar(df: pd.DataFrame, hoja: str) -> ResultadoLectura:
    advertencias = []
    for c in COLUMNAS_SAC:
        if c not in df.columns:
            df[c] = None
            if c not in ("FECHA_PAGO", "TIPO_DE_CAMBIO"):
                advertencias.append(f"La columna {c} no viene en el archivo.")
    df = df[COLUMNAS_SAC].copy()
    for c in _COLUMNAS_NUMERICAS:
        df[c] = _a_numero(df[c])
    for c in _COLUMNAS_FECHA:
        df[c] = _a_fecha(df[c])
    df["CUENTA_GESTION"] = _a_numero(df["CUENTA_GESTION"]).astype("Int64")
    for c in ("DESC_UNID_PRESTACION_SERVICIO", "DESCRIPCION_UNIDAD_FACTURACION", "CLIENTE", "DETALLE"):
        df[c] = df[c].astype("string").str.strip()
    sin_monto = int(df["MONTO"].isna().sum())
    if sin_monto:
        advertencias.append(f"{sin_monto} filas sin monto numérico fueron omitidas.")
    df = df[df["MONTO"].notna() & df["UNIDAD_PRESTACION_SERVICIO"].notna()].reset_index(drop=True)
    df["UNIDAD_PRESTACION_SERVICIO"] = df["UNIDAD_PRESTACION_SERVICIO"].astype(int)
    df["MES"] = df["MES"].astype("Int64")
    # Año contable: el del período de emisión (igual que la columna AÑO del libro: =AÑO(FECHA_EMISION)).
    df["ANIO"] = df["FECHA_EMISION"].dt.year.astype("Int64")
    return ResultadoLectura(datos=df, hoja=hoja, filas_leidas=len(df), advertencias=advertencias)
