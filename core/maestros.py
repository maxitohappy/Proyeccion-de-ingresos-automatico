"""Tablas maestras: plan de cuentas, conceptos y unidades aeroportuarias.

Los maestros institucionales no se publican en el repositorio: se importan desde el
libro "Informe de Gestión" del usuario (hojas PlanDeCuentas, Conceptos, asociación y
matriz) y quedan guardados solo en su equipo/navegador.
"""
from __future__ import annotations

import io
import unicodedata
from dataclasses import dataclass

import pandas as pd

from . import almacenamiento as alm
from .ingest import _a_bytes, normalizar_encabezado

ARCHIVO_PLAN = "maestros/plan_cuentas.csv"
ARCHIVO_CONCEPTOS = "maestros/conceptos.csv"
ARCHIVO_UNIDADES = "maestros/unidades.csv"

# Mapeo Categoría -> Concepto (hoja "Conceptos"); no es información sensible.
CONCEPTOS_DEFECTO = pd.DataFrame(
    [
        ("Ingresos DSO", "Otros Ingresos"),
        ("Autorizaciones y Certificaciones", "Otros Ingresos"),
        ("Combustible", "Combustible"),
        ("Concesiones", "Concesiones"),
        ("Derecho Embarque", "Derecho de Embarque"),
        ("Meteorología", "Otros Ingresos"),
        ("Otros Ingresos", "Otros Ingresos"),
        ("Publicaciones", "Otros Ingresos"),
        ("Reglamentos", "Otros Ingresos"),
        ("Tasas", "Tasas"),
        ("Ticas", "Otros Ingresos"),
        ("Habilitación y Licencias", "Otros Ingresos"),
        ("Certificaciones y Autorizaciones", "Otros Ingresos"),
        ("Servicios Aeroportuarios", "Otros Ingresos"),
    ],
    columns=["categoria", "concepto"],
)

COLUMNAS_PLAN = [
    "cuenta_gestion", "cuenta_presupuestaria", "producto", "subproducto",
    "denominacion", "categoria_tarifaria", "categoria", "tipo_ingreso",
]
COLUMNAS_UNIDADES = [
    "unidad", "nombre", "categoria", "region", "considerar_pax", "considerar_op",
    "unidad_proyeccion", "asociar_a",
    "tasa_aterrizaje", "tasa_estacionamiento", "tasa_iluminacion", "tasa_ils",
    "tasa_servicio_ruta", "derechos_embarque", "derechos_combustible",
]
_CATEGORIA_TEXTO = {"primera": 1, "segunda": 2, "tercera": 3, "cuarta": 4}


@dataclass
class Maestros:
    plan: pd.DataFrame
    conceptos: pd.DataFrame
    unidades: pd.DataFrame

    @property
    def completos(self) -> bool:
        return not self.plan.empty and not self.unidades.empty


def _sin_tildes(texto: str) -> str:
    return unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().lower().strip()


def _si(valor) -> bool:
    return _sin_tildes(valor) in ("si", "s", "true", "1", "x")


def cargar() -> Maestros:
    plan = alm.leer_csv(ARCHIVO_PLAN, dtype={"cuenta_gestion": "Int64"})
    conceptos = alm.leer_csv(ARCHIVO_CONCEPTOS)
    unidades = alm.leer_csv(ARCHIVO_UNIDADES)
    return Maestros(
        plan=plan if plan is not None else pd.DataFrame(columns=COLUMNAS_PLAN),
        conceptos=conceptos if conceptos is not None else CONCEPTOS_DEFECTO.copy(),
        unidades=_tipar_unidades(unidades) if unidades is not None else pd.DataFrame(columns=COLUMNAS_UNIDADES),
    )


def guardar(m: Maestros) -> None:
    alm.escribir_csv(ARCHIVO_PLAN, m.plan)
    alm.escribir_csv(ARCHIVO_CONCEPTOS, m.conceptos)
    alm.escribir_csv(ARCHIVO_UNIDADES, m.unidades)


def _tipar_unidades(u: pd.DataFrame) -> pd.DataFrame:
    u = u.copy()
    for c in COLUMNAS_UNIDADES:
        if c not in u.columns:
            u[c] = None
    u["unidad"] = pd.to_numeric(u["unidad"], errors="coerce").astype("Int64")
    u["asociar_a"] = pd.to_numeric(u["asociar_a"], errors="coerce").astype("Int64")
    u["categoria"] = pd.to_numeric(u["categoria"], errors="coerce").astype("Int64")
    u["region"] = pd.to_numeric(u["region"], errors="coerce").astype("Int64")
    for c in COLUMNAS_UNIDADES[4:7] + COLUMNAS_UNIDADES[8:]:
        u[c] = u[c].map(lambda v: _si(v) if isinstance(v, str) else bool(v) if pd.notna(v) else False)
    return u[COLUMNAS_UNIDADES].dropna(subset=["unidad"]).drop_duplicates("unidad").reset_index(drop=True)


def _hoja(wb, *nombres):
    disponibles = {_sin_tildes(n): n for n in wb.sheetnames}
    for n in nombres:
        if _sin_tildes(n) in disponibles:
            return wb[disponibles[_sin_tildes(n)]]
    return None


def _filas(ws, max_col):
    return [list(f) for f in ws.iter_rows(values_only=True, max_col=max_col) if f and any(v is not None for v in f)]


def importar_desde_libro(origen) -> tuple[Maestros, list[str]]:
    """Construye los maestros desde el libro Informe de Gestión (no guarda)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(_a_bytes(origen)), read_only=True, data_only=True)
    avisos = []
    try:
        plan = _importar_plan(_hoja(wb, "PlanDeCuentas"), avisos)
        conceptos = _importar_conceptos(_hoja(wb, "Conceptos"), avisos)
        unidades = _importar_unidades(_hoja(wb, "asociación", "asociacion"), _hoja(wb, "matriz"), avisos)
    finally:
        wb.close()
    return Maestros(plan=plan, conceptos=conceptos, unidades=unidades), avisos


def _importar_plan(ws, avisos) -> pd.DataFrame:
    if ws is None:
        avisos.append("No se encontró la hoja 'PlanDeCuentas'.")
        return pd.DataFrame(columns=COLUMNAS_PLAN)
    filas = _filas(ws, 9)
    # Columnas B–I: Cuenta de Gestión, Cta. Presupuestaria, Producto, Subproducto,
    # Denominación, Categoría (tarifaria), Categoría, Tipo de Ingreso (como los BUSCARV del libro).
    df = pd.DataFrame([f[1:9] for f in filas[1:]], columns=COLUMNAS_PLAN)
    df["cuenta_gestion"] = pd.to_numeric(df["cuenta_gestion"], errors="coerce").astype("Int64")
    df = df.dropna(subset=["cuenta_gestion"])
    for c in COLUMNAS_PLAN[1:]:
        df[c] = df[c].map(lambda v: str(v).strip() if v is not None else None)
    avisos.append(f"Plan de cuentas: {len(df)} filas ({df['cuenta_gestion'].nunique()} cuentas de gestión).")
    return df.reset_index(drop=True)


def _importar_conceptos(ws, avisos) -> pd.DataFrame:
    if ws is None:
        avisos.append("No se encontró la hoja 'Conceptos'; se usa la tabla por defecto.")
        return CONCEPTOS_DEFECTO.copy()
    filas = [f[:2] for f in _filas(ws, 2)[1:] if f[0] and f[1]]
    return pd.DataFrame([(str(a).strip(), str(b).strip()) for a, b in filas], columns=["categoria", "concepto"])


def _importar_unidades(ws_asoc, ws_matriz, avisos) -> pd.DataFrame:
    if ws_asoc is None:
        avisos.append("No se encontró la hoja 'asociación'.")
        return pd.DataFrame(columns=COLUMNAS_UNIDADES)
    filas = _filas(ws_asoc, 8)
    enc = [normalizar_encabezado(h) for h in filas[0]]
    asoc = pd.DataFrame(filas[1:], columns=enc)
    u = pd.DataFrame({
        "unidad": pd.to_numeric(asoc.iloc[:, 0], errors="coerce"),
        "nombre": asoc.iloc[:, 1].map(lambda v: " ".join(str(v).split()) if v else None),
        "categoria": pd.to_numeric(asoc.iloc[:, 2], errors="coerce"),
        "region": pd.to_numeric(asoc.iloc[:, 3], errors="coerce"),
        "considerar_pax": asoc.iloc[:, 4].map(_si),
        "considerar_op": asoc.iloc[:, 5].map(_si),
        "unidad_proyeccion": asoc.iloc[:, 6].map(_si),
        "asociar_a": pd.to_numeric(asoc.iloc[:, 7], errors="coerce"),
    })
    if ws_matriz is not None:
        fm = _filas(ws_matriz, 10)
        m = pd.DataFrame(fm[1:], columns=[
            "unidad", "nombre_matriz", "categoria_texto", "tasa_aterrizaje", "tasa_estacionamiento",
            "tasa_iluminacion", "tasa_ils", "tasa_servicio_ruta", "derechos_embarque", "derechos_combustible",
        ])
        m["unidad"] = pd.to_numeric(m["unidad"], errors="coerce")
        m["nombre_matriz"] = m["nombre_matriz"].map(lambda v: " ".join(str(v).split()) if v else None)
        u = u.merge(m, on="unidad", how="outer")
        u["nombre"] = u["nombre_matriz"].fillna(u["nombre"])
        cat_texto = u["categoria_texto"].map(lambda v: _CATEGORIA_TEXTO.get(_sin_tildes(v)) if v else None)
        u["categoria"] = u["categoria"].fillna(cat_texto)
        u = u.drop(columns=["nombre_matriz", "categoria_texto"])
    u = _tipar_unidades(u)
    avisos.append(f"Unidades: {len(u)} ({int(u['asociar_a'].notna().sum())} asociadas a otra unidad).")
    return u


def unidad_grupo(unidades: pd.DataFrame) -> dict[int, int]:
    """Unidad -> unidad a la que se asocian sus ingresos (resuelve cadenas)."""
    padre = {int(r.unidad): int(r.asociar_a) for r in unidades.itertuples() if pd.notna(r.asociar_a)}
    resultado = {}
    for u in unidades["unidad"].dropna().astype(int):
        actual, visitadas = u, set()
        while actual in padre and actual not in visitadas:
            visitadas.add(actual)
            actual = padre[actual]
        resultado[u] = actual
    return resultado


def nombres_unidades(unidades: pd.DataFrame) -> dict[int, str]:
    return {int(r.unidad): str(r.nombre) for r in unidades.itertuples() if pd.notna(r.unidad)}
